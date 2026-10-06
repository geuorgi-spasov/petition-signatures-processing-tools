"""Tests for split_signatures_into_folders.py."""
from __future__ import annotations

import pandas as pd
import pytest

from split_signatures_into_folders import (
    BOTTOM_MARGIN_CM,
    CSV_ENCODINGS,
    CSV_SEPARATORS,
    DEFAULT_COLUMN_WIDTHS_CM,
    FOOTER_FONT_SIZE_PT,
    FOOTER_LINES,
    FOOTER_LINE_SPACING,
    FOOTER_SPACE_BEFORE_PT,
    INPUT_CSV,
    PAGE_HEIGHT_CM,
    PAGE_SAFETY_MARGIN_CM,
    POINTS_PER_CM,
    ROW_HEIGHT_CM,
    TOP_MARGIN_CM,
    footer_block_cm,
    LEFT_MARGIN_CM,
    OUTPUT_DOCX_FOLDER,
    PAGE_WIDTH_CM,
    RIGHT_MARGIN_CM,
    ROWS_PER_FILE,
    ROWS_PER_PAGE,
    apply_full_table_borders,
    build_arg_parser,
    build_folder_document,
    fits_on_the_page,
    in_blocks_of,
    main,
    max_rows_per_page,
    read_signatures_csv,
    scale_column_widths,
)

AVAILABLE_WIDTH_CM = PAGE_WIDTH_CM - LEFT_MARGIN_CM - RIGHT_MARGIN_CM


# ---------------------------------------------------------------------------
# scale_column_widths
# ---------------------------------------------------------------------------

class TestScaleColumnWidths:
    def test_six_columns_fills_available_width(self):
        widths = scale_column_widths(6)
        assert len(widths) == 6
        assert sum(widths) == pytest.approx(AVAILABLE_WIDTH_CM)

    def test_six_columns_preserves_default_proportions(self):
        widths = scale_column_widths(6)
        ratios = [w / d for w, d in zip(widths, DEFAULT_COLUMN_WIDTHS_CM)]
        for r in ratios[1:]:
            assert r == pytest.approx(ratios[0])

    def test_fewer_columns_uses_first_n_defaults(self):
        widths = scale_column_widths(3)
        assert len(widths) == 3
        assert sum(widths) == pytest.approx(AVAILABLE_WIDTH_CM)

    def test_more_columns_than_defaults_distributes_evenly(self):
        widths = scale_column_widths(8)
        assert len(widths) == 8
        assert all(w == pytest.approx(widths[0]) for w in widths)
        assert sum(widths) == pytest.approx(AVAILABLE_WIDTH_CM)

    def test_one_column_takes_full_width(self):
        widths = scale_column_widths(1)
        assert widths == pytest.approx([AVAILABLE_WIDTH_CM])


# ---------------------------------------------------------------------------
# read_signatures_csv
# ---------------------------------------------------------------------------

class TestReadSignaturesCsv:
    def _write(self, path, content, encoding="utf-8"):
        path.write_text(content, encoding=encoding)

    def test_reads_comma_separated_with_headers(self, tmp_path):
        csv = tmp_path / "sigs.csv"
        self._write(csv, "id,name,city\n1,Иван,София\n2,Петър,Пловдив\n")
        df = read_signatures_csv(str(csv))
        assert list(df.columns) == ["id", "name", "city"]
        assert len(df) == 2

    def test_reads_semicolon_separated(self, tmp_path):
        csv = tmp_path / "sigs.csv"
        self._write(csv, "id;name;city\n1;Иван;София\n2;Петър;Пловдив\n")
        df = read_signatures_csv(str(csv))
        assert len(df) == 2
        assert df.iloc[0]["city"] == "София"

    def test_raises_on_single_column_file(self, tmp_path):
        csv = tmp_path / "broken.csv"
        self._write(csv, "only_one_column\nrow1\nrow2\n")
        with pytest.raises(ValueError):
            read_signatures_csv(str(csv))


# ---------------------------------------------------------------------------
# build_folder_document — smoke test that exercises the whole pipeline
# ---------------------------------------------------------------------------

class TestBuildFolderDocument:
    def _make_folder_rows(self, n_rows):
        return pd.DataFrame(
            {
                "id": range(1, n_rows + 1),
                "name": [f"Name {i}" for i in range(1, n_rows + 1)],
                "city": ["София"] * n_rows,
                "address": ["Address line " * 3] * n_rows,
                "date": ["2024-01-01"] * n_rows,
                "signature": ["OK"] * n_rows,
            }
        )

    def test_returns_document_and_correct_page_count(self):
        folder_rows = self._make_folder_rows(25)  # 25 rows = 3 pages at 10 per page
        header = folder_rows.columns.tolist()
        widths = scale_column_widths(len(header))

        doc, pages = build_folder_document(1, folder_rows, header, widths)

        assert pages == 3
        # One table per page
        assert len(doc.tables) == 3

    def test_single_page_for_small_chunk(self):
        folder_rows = self._make_folder_rows(5)
        header = folder_rows.columns.tolist()
        widths = scale_column_widths(len(header))

        doc, pages = build_folder_document(1, folder_rows, header, widths)
        assert pages == 1
        assert len(doc.tables) == 1

    def test_full_folder_size(self):
        folder_rows = self._make_folder_rows(1000)
        header = folder_rows.columns.tolist()
        widths = scale_column_widths(len(header))

        doc, pages = build_folder_document(1, folder_rows, header, widths)
        assert pages == 100  # 1000 / 10
        assert len(doc.tables) == 100


# ---------------------------------------------------------------------------
# The page's height is the program's to decide, not the renderer's
# ---------------------------------------------------------------------------

class TestNothingOnThePageIsLeftToTheRenderer:
    """The bug these exist for.

    The table was pinned to the millimetre with EXACTLY row heights while
    the footer's spacing was left unset, so whatever opened the document
    chose it. LibreOffice chose small enough to fit; Word chose larger,
    the footer moved onto a page of its own, and a 100-page folder opened
    as 199. The declared page count is printed in that very footer, so
    the one thing this toolkit may not get wrong was decided elsewhere.
    """

    def _page(self):
        rows = pd.DataFrame({name: ["x"] * 20 for name in
                             ["Номер", "Име", "Фамилия", "Имейл", "Дата", "Шифър"]})
        doc, _ = build_folder_document(1, rows, list(rows.columns),
                                      scale_column_widths(6))
        return doc

    def test_every_paragraph_has_a_line_height_of_its_own(self):
        # Including the one that only carries a page break: its line box
        # sits on the page it ends, so its height is part of the page.
        doc = self._page()
        assert doc.paragraphs, "no body paragraphs to check"
        for paragraph in doc.paragraphs:
            spacing = paragraph.paragraph_format
            assert spacing.line_spacing is not None, paragraph.text[:30]
            assert spacing.space_after.pt == 0, paragraph.text[:30]

    def test_the_line_height_is_the_footer_font_times_its_spacing(self):
        # Compared in points: a Length is stored in EMU, and the same
        # 14.4 pt built twice can land one EMU apart.
        doc = self._page()
        expected = FOOTER_FONT_SIZE_PT * FOOTER_LINE_SPACING
        for paragraph in doc.paragraphs:
            assert paragraph.paragraph_format.line_spacing.pt == (
                pytest.approx(expected)
            )

    def test_the_footer_does_not_open_with_a_blank_line(self):
        # It used to, for the gap above it. The gap is space_before now,
        # which is a number the arithmetic can count.
        doc = self._page()
        footer = doc.paragraphs[0]
        assert footer._p.xml.count("<w:br") == 1, "one break, between the runs"
        assert footer.paragraph_format.space_before.pt == (
            pytest.approx(FOOTER_SPACE_BEFORE_PT)
        )

    def test_the_page_break_paragraph_is_pinned_too(self):
        doc = self._page()
        breaks = [p for p in doc.paragraphs if 'type="page"' in p._p.xml]
        assert breaks, "a 20-row folder should carry one page break"
        for paragraph in breaks:
            assert paragraph.paragraph_format.line_spacing is not None
            assert paragraph.paragraph_format.space_after.pt == 0


class TestFooterBlockIsDerived:
    def test_it_is_the_lines_plus_the_gap_above_them(self):
        points = (FOOTER_SPACE_BEFORE_PT
                  + FOOTER_LINES * FOOTER_FONT_SIZE_PT * FOOTER_LINE_SPACING)
        assert footer_block_cm() == pytest.approx(points / POINTS_PER_CM)

    def test_it_counts_the_page_break_paragraph_as_one_of_its_lines(self):
        # Two printed lines and the paragraph carrying the break.
        assert FOOTER_LINES == 3

    def test_the_default_page_keeps_the_safety_margin_in_hand(self):
        usable = PAGE_HEIGHT_CM - TOP_MARGIN_CM - BOTTOM_MARGIN_CM
        table = (ROWS_PER_PAGE + 1) * ROW_HEIGHT_CM
        spare = usable - table - footer_block_cm()
        assert spare >= PAGE_SAFETY_MARGIN_CM, (
            f"only {spare:.2f} cm spare, less than the "
            f"{PAGE_SAFETY_MARGIN_CM} cm meant to be left over"
        )

    def test_one_more_row_would_eat_the_margin(self):
        # Which is why the cap is where it is.
        assert not fits_on_the_page(ROWS_PER_PAGE + 1)


# ---------------------------------------------------------------------------
# Encodings
# ---------------------------------------------------------------------------

BULGARIAN_ROWS = "Номер\tИме\tФамилия\n1\tИван\tИванов\n2\tМария\tМаринова\n"


def _write(path, text, encoding):
    path.write_bytes(text.encode(encoding))
    return str(path)


class TestEncodings:
    @pytest.mark.parametrize("encoding", ["utf-8", "utf-8-sig", "cp1251"])
    def test_bulgarian_names_survive_the_round_trip(self, tmp_path, encoding):
        path = _write(tmp_path / "signatures.csv", BULGARIAN_ROWS, encoding)
        df = read_signatures_csv(path)
        assert list(df.columns) == ["Номер", "Име", "Фамилия"]
        assert list(df.iloc[0])[1:] == ["Иван", "Иванов"]

    def test_the_last_resort_is_last(self):
        # latin-1 decodes any byte at all, so anything after it could never
        # be reached - and anything before it must be able to fail.
        assert CSV_ENCODINGS[-1] == "latin-1"
        assert bytes(range(256)).decode("latin-1")
        for earlier in CSV_ENCODINGS[:-1]:
            with pytest.raises(UnicodeDecodeError):
                bytes(range(256)).decode(earlier)

    def test_every_encoding_in_the_list_is_reachable(self):
        # An encoding is only worth listing if the ones before it reject
        # the bytes it is there to catch. The last resort is exempt - its
        # whole job is to catch what nothing else did.
        middle = CSV_ENCODINGS[1:-1]
        assert middle, "nothing between utf-8 and the last resort to check"
        for position, encoding in enumerate(middle, start=1):
            sample = BULGARIAN_ROWS.encode(encoding)
            assert any(
                not _decodes(sample, earlier)
                for earlier in CSV_ENCODINGS[:position]
            ), f"{encoding} is shadowed by the encodings before it"

    @pytest.mark.parametrize("shadowed", ["cp866", "cp1252", "iso-8859-1"])
    def test_the_encodings_left_out_really_are_unreachable(self, shadowed):
        # These were in the list once. Each is decoded without error by an
        # encoding that comes before it, so it could never have been used.
        sample = BULGARIAN_ROWS.encode("cp1251")
        assert _decodes(sample, "cp1251")
        assert _decodes(bytes(range(256)), "latin-1")
        assert shadowed not in CSV_ENCODINGS

    def test_a_missing_file_is_reported_as_missing(self, tmp_path):
        # Not as a parse failure: there is nothing to parse, and listing
        # every separator and encoding "tried" would be a lie. main()
        # has always had a handler for this - it just could not fire.
        with pytest.raises(FileNotFoundError):
            read_signatures_csv(str(tmp_path / "does_not_exist.csv"))

    def test_an_unreadable_file_says_what_was_tried(self, tmp_path):
        path = tmp_path / "one_column.csv"
        path.write_text("just_one_column\nstill_one\n", encoding="utf-8")
        with pytest.raises(ValueError) as error:
            read_signatures_csv(str(path))
        assert "Tried separators" in str(error.value)


def _decodes(data: bytes, encoding: str) -> bool:
    try:
        data.decode(encoding)
        return True
    except UnicodeDecodeError:
        return False


# ---------------------------------------------------------------------------
# Separators
# ---------------------------------------------------------------------------

class TestSeparators:
    def test_a_comma_inside_a_field_does_not_beat_the_real_separator(
        self, tmp_path
    ):
        # The export this script reads has no field with a comma in it -
        # "Имейл адрес" is an email address, not a postal one. This is a
        # hypothetical third column, standing for any future field that
        # does: the file is tab-separated, so the tab must win or every
        # field shifts one column left.
        path = tmp_path / "signatures.csv"
        path.write_text(
            "Номер\tИме\tБележка\n"
            "1\tИван Иванов\tпърва, втора\n"
            "2\tМария Маринова\tтрета, четвърта\n",
            encoding="utf-8",
        )
        df = read_signatures_csv(str(path))
        assert list(df.columns) == ["Номер", "Име", "Бележка"]
        assert df.iloc[0]["Бележка"] == "първа, втора"

    def test_a_genuine_comma_file_still_works(self, tmp_path):
        # Ordering the comma last must not stop it being found when it
        # really is the separator.
        path = tmp_path / "signatures.csv"
        path.write_text("Номер,Име\n1,Иван\n2,Мария\n", encoding="utf-8")
        df = read_signatures_csv(str(path))
        assert list(df.columns) == ["Номер", "Име"]

    @pytest.mark.parametrize("inside", [";", "|"])
    def test_a_separator_inside_a_field_does_not_beat_the_tab(
        self, tmp_path, inside
    ):
        # Both splits give the same number of columns, so no count can
        # choose between them and the order is what decides. Splitting on
        # the ';' here would shift every field one column left.
        path = tmp_path / "signatures.csv"
        path.write_text(
            f"Номер\tИме\tБележка\n"
            f"1\tИван Иванов\tпърва{inside}втора\n"
            f"2\tМария Маринова\tтрета{inside}четвърта\n",
            encoding="utf-8",
        )
        df = read_signatures_csv(str(path))
        assert list(df.columns) == ["Номер", "Име", "Бележка"]
        assert df.iloc[0]["Бележка"] == f"първа{inside}втора"

    def test_the_tab_is_tried_first_and_the_comma_last(self):
        # A separator cannot fail the way an encoding can, so the order is
        # the only guard: least likely inside a field first, most likely
        # last. A tab is not a character anyone types mid-word; a comma is.
        assert CSV_SEPARATORS[0] == "\t"
        assert CSV_SEPARATORS[-1] == ","
        assert set(CSV_SEPARATORS) == {"\t", "|", ";", ","}


# ---------------------------------------------------------------------------
# Table borders
# ---------------------------------------------------------------------------

class TestTableBorders:
    def _table(self):
        from docx import Document
        doc = Document()
        table = doc.add_table(rows=3, cols=4)
        apply_full_table_borders(table)
        return table

    def test_the_table_gets_an_outer_frame_and_inner_gridlines(self):
        xml = self._table()._tbl.xml
        assert "<w:tblBorders>" in xml
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            assert f"<w:{edge} " in xml, f"table edge {edge} missing"

    def test_every_cell_gets_its_own_four_sides(self):
        # Belt and braces: some viewers honour only the cell-level borders.
        table = self._table()
        for row in table.rows:
            for cell in row.cells:
                xml = cell._tc.xml
                assert "<w:tcBorders>" in xml
                assert xml.count("w:val=\"single\"") >= 4

    def test_the_line_is_thin_solid_and_black(self):
        xml = self._table()._tbl.xml
        assert 'w:val="single"' in xml   # solid, not dashed
        assert 'w:sz="4"' in xml         # eighths of a point -> 0.5 pt
        assert 'w:color="000000"' in xml


# ---------------------------------------------------------------------------
# Cutting the rows into blocks
# ---------------------------------------------------------------------------

class TestInBlocksOf:
    def _rows(self, n):
        return pd.DataFrame({"a": range(n)})

    @pytest.mark.parametrize("total,size,expected", [
        (10, 10, [10]),          # exactly one full block
        (11, 10, [10, 1]),       # the ragged last block
        (9, 10, [9]),            # less than one block
        (25, 10, [10, 10, 5]),
        (1000, 10, [10] * 100),  # one real submission folder
        (0, 10, []),             # nothing at all
    ])
    def test_the_last_block_is_whatever_is_left(self, total, size, expected):
        blocks = list(in_blocks_of(self._rows(total), size))
        assert [len(b) for b in blocks] == expected

    def test_no_row_is_lost_or_repeated(self):
        rows = self._rows(1001)
        seen = pd.concat(list(in_blocks_of(rows, 10)))
        assert list(seen["a"]) == list(range(1001))


# ---------------------------------------------------------------------------
# The command line
# ---------------------------------------------------------------------------

class TestCommandLine:
    def test_the_defaults_are_the_constants(self):
        args = build_arg_parser().parse_args([])
        assert args.input == INPUT_CSV
        assert args.output == OUTPUT_DOCX_FOLDER
        assert args.rows_per_file == ROWS_PER_FILE
        assert args.rows_per_page == ROWS_PER_PAGE

    def test_a_dry_run_writes_nothing(self, tmp_path, monkeypatch, capsys):
        csv = tmp_path / "s.csv"
        csv.write_text("Номер\tИме\tФамилия\n1\tИван\tИванов\n", encoding="utf-8")
        monkeypatch.chdir(tmp_path)
        code = main(["--input", str(csv), "--output", "out", "--dry-run"])
        assert code == 0
        assert not (tmp_path / "out").exists()
        assert "nothing written" in capsys.readouterr().out

    def test_a_page_that_cannot_fit_is_refused(self, capsys):
        # The table would run off the page and the declared page count
        # would stop matching the rendered one - the one thing this
        # toolkit may not get wrong.
        code = main(["--rows-per-page", str(max_rows_per_page() + 1),
                     "--dry-run"])
        assert code == 1
        assert "do not fit" in capsys.readouterr().out

    def test_the_default_page_fits(self):
        assert fits_on_the_page(ROWS_PER_PAGE)
        assert ROWS_PER_PAGE <= max_rows_per_page()

    def test_limit_shortens_the_run(self, tmp_path, capsys):
        csv = tmp_path / "s.csv"
        rows = "".join(f"{i}\tИван\tИванов\n" for i in range(1, 51))
        csv.write_text("Номер\tИме\tФамилия\n" + rows, encoding="utf-8")
        main(["--input", str(csv), "--limit", "20", "--dry-run"])
        assert "20 signatures" in capsys.readouterr().out
